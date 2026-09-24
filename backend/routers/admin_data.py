"""Admin APIs for configured languages and small clinic knowledge documents."""
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy.orm import Session

from database.connection import get_db
from database.models import KnowledgeDocument, LanguageSetting, VoiceProfile
from config import settings
from security import require_admin, supported_language_map

router = APIRouter(prefix="/api/admin", tags=["Admin data"], dependencies=[Depends(require_admin)])
public_router = APIRouter(prefix="/api", tags=["Public settings"])

LANGUAGES = {
    "en": ("English", "en-US"),
    "ta": ("Tamil", "ta"),
    "zh": ("Mandarin", "zh-CN"),
    "ms": ("Malay", "ms-MY"),
}
MAX_DOCUMENT_BYTES = 2 * 1024 * 1024


def effective_languages(db: Session):
    stt = supported_language_map()
    configured_tts = {item.strip().lower() for item in settings.TESTED_TTS_LANGUAGES.split(",") if item.strip()}
    rows = db.query(LanguageSetting).order_by(LanguageSetting.id).all()
    return [
        {"code": row.code, "name": row.name, "enabled": bool(row.enabled),
         "supported": row.code in stt and row.code in configured_tts,
         "deepgramLanguage": stt.get(row.code)}
        for row in rows
    ]


@public_router.get("/public-settings")
def public_settings(db: Session = Depends(get_db)):
    active_voice = db.query(VoiceProfile).filter_by(is_active=True).first()
    providers_ready = bool(active_voice and settings.DEEPGRAM_API_KEY and settings.ELEVENLABS_API_KEY and settings.OPENAI_API_KEY)
    return {"languages": [x for x in effective_languages(db) if providers_ready and x["enabled"] and x["supported"]],
            "voiceAvailable": providers_ready}


class LanguageUpdate(BaseModel):
    enabled: bool


@router.get("/languages")
def list_languages(db: Session = Depends(get_db)):
    return {"languages": effective_languages(db), "supportedCodes": sorted(set(supported_language_map()) & {
        item.strip().lower() for item in settings.TESTED_TTS_LANGUAGES.split(",") if item.strip()
    })}


@router.put("/languages/{code}")
def update_language(code: str, update: LanguageUpdate, db: Session = Depends(get_db)):
    row = db.query(LanguageSetting).filter_by(code=code.lower()).first()
    if not row:
        raise HTTPException(404, "Language not found")
    available = {**supported_language_map()}
    tts = {item.strip().lower() for item in settings.TESTED_TTS_LANGUAGES.split(",") if item.strip()}
    if row.code == "en" and not update.enabled:
        raise HTTPException(400, "English must remain enabled for the spoken language-selection prompt")
    if update.enabled and (row.code not in available or row.code not in tts):
        raise HTTPException(400, "Language cannot be enabled until STT and voice output are marked tested")
    row.enabled = update.enabled
    db.commit()
    return {"language": next(x for x in effective_languages(db) if x["code"] == row.code)}


@router.get("/knowledge")
def list_documents(db: Session = Depends(get_db)):
    return [{"id": d.id, "filename": d.filename, "contentType": d.content_type,
             "sizeBytes": d.size_bytes, "status": d.status, "error": d.error,
             "createdAt": d.created_at.isoformat()} for d in db.query(KnowledgeDocument).order_by(KnowledgeDocument.created_at.desc()).all()]


def _extract(filename: str, content_type: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"} and content_type in {"text/plain", "text/markdown", "application/octet-stream"}:
        try:
            return payload.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(400, "Text documents must be UTF-8 encoded") from exc
    if suffix == ".pdf" and content_type == "application/pdf":
        try:
            from pypdf import PdfReader
            import io
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(payload)).pages)
        except Exception as exc:
            raise HTTPException(400, f"Could not read PDF: {exc}") from exc
        if not text.strip():
            raise HTTPException(422, "PDF has no extractable text; scanned/image-only PDFs are not supported")
        return text
    if suffix == ".docx" and content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            from docx import Document
            import io
            return "\n".join(p.text for p in Document(io.BytesIO(payload)).paragraphs)
        except Exception as exc:
            raise HTTPException(400, f"Could not read DOCX: {exc}") from exc
    raise HTTPException(415, "Supported formats: UTF-8 .txt/.md, text-based .pdf, and .docx")


@router.post("/knowledge")
async def upload_document(file: UploadFile = File(...), db: Session = Depends(get_db)):
    filename = Path(file.filename or "").name
    if not filename or len(filename) > 255:
        raise HTTPException(400, "Invalid filename")
    payload = await file.read(MAX_DOCUMENT_BYTES + 1)
    if not payload or len(payload) > MAX_DOCUMENT_BYTES:
        raise HTTPException(413, "Document must be between 1 byte and 2 MiB")
    try:
        text = _extract(filename, file.content_type or "", payload)
    except UnicodeDecodeError as exc:
        raise HTTPException(400, "Text documents must be UTF-8 encoded") from exc
    if not text.strip() or len(text) > 500_000:
        raise HTTPException(422, "Extracted document text must be non-empty and at most 500,000 characters")
    record = KnowledgeDocument(filename=filename, content_type=file.content_type or "application/octet-stream",
        size_bytes=len(payload), extracted_text=text, original_bytes=payload, status="ready")
    db.add(record)
    db.commit()
    db.refresh(record)
    return {"id": record.id, "filename": record.filename, "status": record.status, "sizeBytes": record.size_bytes}


@router.delete("/knowledge/{document_id}")
def delete_document(document_id: int, db: Session = Depends(get_db)):
    record = db.query(KnowledgeDocument).filter_by(id=document_id).first()
    if not record:
        raise HTTPException(404, "Document not found")
    db.delete(record)
    db.commit()
    return {"success": True}
